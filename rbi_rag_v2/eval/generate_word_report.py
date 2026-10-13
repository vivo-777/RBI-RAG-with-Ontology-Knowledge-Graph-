import os
import re
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_COLOR_INDEX
except ImportError:
    print("Please install python-docx using: pip install python-docx")
    exit(1)

def convert_md_to_docx(md_path, docx_path):
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found.")
        return

    doc = Document()
    doc.add_heading('RBI RAG Pipeline - Comparative Evaluation Results', 0)

    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    # Split by query blocks (using ### X. Question:)
    blocks = re.split(r'(### \d+\. Question:.*)', content)
    
    # The first block is the summary/header
    summary = blocks[0].strip()
    if summary:
        doc.add_heading('Aggregate Metrics', level=1)
        doc.add_paragraph(summary)

    # Process each query block
    for i in range(1, len(blocks), 2):
        question_header = blocks[i].strip()
        body = blocks[i+1].strip()

        # Add question heading
        p = doc.add_heading(level=2)
        run = p.add_run(question_header.replace('### ', ''))
        run.bold = True
        run.font.color.rgb = RGBColor(0, 51, 102)

        # Parse body
        sections = re.split(r'(\*\*Answer:\*\*|\*\*Top Sources Consulted:\*\*|> \*System Metrics.*?$)', body, flags=re.MULTILINE)
        
        current_section = None
        for sec in sections:
            sec = sec.strip()
            if not sec:
                continue
                
            if sec == '**Answer:**':
                current_section = 'answer'
                doc.add_heading('Pipeline Answer:', level=3)
            elif sec == '**Top Sources Consulted:**':
                current_section = 'sources'
                doc.add_heading('Sources:', level=3)
            elif sec.startswith('> *System Metrics'):
                current_section = 'metrics'
                p = doc.add_paragraph()
                run = p.add_run(sec.replace('> *', '').replace('*', ''))
                run.italic = True
                run.font.color.rgb = RGBColor(128, 128, 128)
            else:
                if current_section == 'answer':
                    doc.add_paragraph(sec)
                elif current_section == 'sources':
                    for line in sec.split('\n'):
                        if line.strip():
                            doc.add_paragraph(line.strip(), style='List Bullet')

        # Add placeholder for Gemini Answer
        doc.add_heading('Gemini Answer:', level=3)
        placeholder = doc.add_paragraph()
        run = placeholder.add_run('[User to insert Gemini answer here]')
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        
        doc.add_page_break()

    doc.save(docx_path)
    print(f"Successfully generated Word document at {docx_path}")

if __name__ == "__main__":
    md_file = os.path.join(os.path.dirname(__file__), "..", "test_results_grok.md")
    docx_file = os.path.join(os.path.dirname(__file__), "..", "test_results_report.docx")
    convert_md_to_docx(md_file, docx_file)
