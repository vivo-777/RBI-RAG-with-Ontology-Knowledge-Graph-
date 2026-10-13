import json
import os
try:
    from docx import Document
except ImportError:
    import sys
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

def create_docx_report():
    input_file = os.path.join(os.path.dirname(__file__), "tailscale_40_results.json")
    output_file = os.path.join(os.path.dirname(__file__), "tailscale_40_results.docx")
    
    with open(input_file, "r", encoding="utf-8") as f:
        results = json.load(f)
        
    doc = Document()
    doc.add_heading("Evaluation Results (Tailscale - Qwen 2.5 7B)", 0)
    
    for i, res in enumerate(results):
        doc.add_heading(f"Query {i+1}:", level=2)
        doc.add_paragraph(res.get("query", ""))
        
        doc.add_heading("Answer:", level=3)
        doc.add_paragraph(res.get("answer", ""))
        
        latency_p = doc.add_paragraph()
        latency_p.add_run(f"Latency: {res.get('latency', 0):.2f}s").italic = True
        
        # Add a separator
        doc.add_paragraph("-" * 40)
        
    doc.save(output_file)
    print(f"Successfully created {output_file}")

if __name__ == "__main__":
    create_docx_report()
