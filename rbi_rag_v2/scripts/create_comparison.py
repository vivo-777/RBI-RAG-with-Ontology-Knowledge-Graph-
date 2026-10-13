import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def parse_md_results(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()

    # Extract aggregate metrics
    metrics_match = re.search(r'## Aggregate Metrics\n\n(.*?)\n\n---', content, re.DOTALL)
    aggregate = metrics_match.group(1).strip() if metrics_match else ""

    # Extract questions
    questions = {}
    blocks = re.split(r'### \d+\. Question:', content)
    for block in blocks[1:]:
        lines = block.strip().split('\n')
        q_text = lines[0].strip()
        
        # Find Answer
        ans_start = block.find('**Answer:**')
        sources_start = block.find('**Top Sources Consulted:**')
        if sources_start == -1:
            sources_start = block.find('**Sources Consulted:**')
            
        metrics_start = block.find('> *System Metrics (FACTUAL):')
        
        answer_text = block[ans_start + 11:sources_start].strip() if ans_start != -1 and sources_start != -1 else ""
        
        # Find Metrics
        sys_metrics = ""
        if metrics_start != -1:
            end_idx = block.find('*', metrics_start + 29)
            if end_idx != -1:
                sys_metrics = block[metrics_start + 29:end_idx].strip()
            else:
                sys_metrics = block[metrics_start + 29:].strip()
                
        questions[q_text] = {
            'answer': answer_text,
            'metrics': sys_metrics
        }
        
    return aggregate, questions

local_agg, local_qs = parse_md_results('test_results.md')
groq_agg, groq_qs = parse_md_results('test_results_grok.md')

doc = Document()

# Title
title = doc.add_heading('RBI Regulatory Assistant: Evaluation Comparison', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph("This document provides a side-by-side comparison of the evaluation results between the Local On-Premise Model (Qwen 2.5 7B) and the Cloud API Model (Groq Llama 3.3 70B).")

# Aggregate Metrics
doc.add_heading('Aggregate Metrics Comparison', level=1)
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Metric'
hdr_cells[1].text = 'Local Model (Qwen 2.5)'
hdr_cells[2].text = 'Cloud API (Groq Llama 3.3)'

def get_metric_val(agg_text, metric_name):
    for line in agg_text.split('\n'):
        if metric_name in line:
            parts = line.split('|')
            if len(parts) >= 3:
                return parts[2].strip()
    return "N/A"

metrics_to_compare = [
    ("Average Faithfulness", "**Average Faithfulness"),
    ("Average Context Recall", "**Average Context Recall**"),
    ("Average Context Precision", "**Average Context Precision**"),
    ("Latency p50", "**Latency p50**"),
]

for display, search_key in metrics_to_compare:
    row_cells = table.add_row().cells
    row_cells[0].text = display
    row_cells[1].text = get_metric_val(local_agg, search_key)
    row_cells[2].text = get_metric_val(groq_agg, search_key)

# Detailed Q&A
doc.add_page_break()
doc.add_heading('Detailed Q&A Comparison', level=1)

all_questions = list(local_qs.keys())
for q in groq_qs.keys():
    if q not in all_questions:
        all_questions.append(q)

for i, q in enumerate(all_questions):
    doc.add_heading(f'Question {i+1}: {q}', level=2)
    
    # Local Model
    p1 = doc.add_paragraph()
    runner = p1.add_run('Local Model (Qwen 2.5 7B)')
    runner.bold = True
    runner.font.color.rgb = RGBColor(0, 102, 204)
    
    local_ans = local_qs.get(q, {}).get('answer', 'N/A')
    local_met = local_qs.get(q, {}).get('metrics', 'N/A')
    
    doc.add_paragraph(local_ans)
    m1 = doc.add_paragraph()
    r1 = m1.add_run(f'Metrics: {local_met}')
    r1.italic = True
    r1.font.size = Pt(9)
    
    # Groq Model
    p2 = doc.add_paragraph()
    runner = p2.add_run('Cloud API (Groq Llama 3.3 70B)')
    runner.bold = True
    runner.font.color.rgb = RGBColor(0, 153, 76)
    
    groq_ans = groq_qs.get(q, {}).get('answer', 'N/A')
    groq_met = groq_qs.get(q, {}).get('metrics', 'N/A')
    
    doc.add_paragraph(groq_ans)
    m2 = doc.add_paragraph()
    r2 = m2.add_run(f'Metrics: {groq_met}')
    r2.italic = True
    r2.font.size = Pt(9)
    
    doc.add_paragraph("_" * 50)

doc.save('Evaluation_Comparison_Report_V2.docx')
print("Comparison document created successfully.")
