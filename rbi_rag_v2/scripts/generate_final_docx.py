import json
import os
from docx import Document

# Load updated JSON
results_path = os.path.join(os.path.dirname(__file__), "tailscale_40_results.json")
with open(results_path, "r", encoding="utf-8") as f:
    results = json.load(f)

docx_path = os.path.join(os.path.dirname(__file__), "tailscale_40_results_final.docx")
doc = Document()
doc.add_heading("Evaluation Results (Tailscale - Qwen 2.5 7B) [FINAL]", 0)

for i, res in enumerate(results):
    doc.add_heading(f"Query {i+1}:", level=2)
    doc.add_paragraph(res.get("query", ""))
    doc.add_heading("Answer:", level=3)
    doc.add_paragraph(res.get("answer", ""))
    latency_p = doc.add_paragraph()
    latency_p.add_run(f"Latency: {res.get('latency', 0):.2f}s").italic = True
    doc.add_paragraph("-" * 40)

doc.save(docx_path)
print(f"Generated new DOCX at {docx_path}")
