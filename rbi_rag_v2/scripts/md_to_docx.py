import os
from docx import Document
from docx.shared import Inches, Pt
import re

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_data = []
    
    for line in lines:
        line = line.strip()
        if not line:
            if not in_table:
                doc.add_paragraph()
            continue
            
        # Headers
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
            
        # Lists
        elif line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            # very basic bold parser
            text = line[2:]
            if '**' in text:
                parts = text.split('**')
                for i, part in enumerate(parts):
                    if i % 2 == 1:
                        p.add_run(part).bold = True
                    else:
                        p.add_run(part)
            else:
                p.add_run(text)
                
        # Blockquotes
        elif line.startswith('> '):
            p = doc.add_paragraph(style='Quote')
            p.add_run(line[2:].strip())
            
        # Tables
        elif line.startswith('|'):
            in_table = True
            if '---' not in line:
                row = [cell.strip().replace('**', '') for cell in line.split('|')[1:-1]]
                table_data.append(row)
        
        else:
            if in_table:
                # End of table
                in_table = False
                if table_data:
                    table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                    table.style = 'Table Grid'
                    for i, row in enumerate(table_data):
                        for j, cell in enumerate(row):
                            table.cell(i, j).text = cell
                    table_data = []
                doc.add_paragraph()
                
            p = doc.add_paragraph()
            # very basic bold parser
            if '**' in line:
                parts = line.split('**')
                for i, part in enumerate(parts):
                    if i % 2 == 1:
                        p.add_run(part).bold = True
                    else:
                        p.add_run(part)
            else:
                p.add_run(line)

    if in_table and table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        table.style = 'Table Grid'
        for i, row in enumerate(table_data):
            for j, cell in enumerate(row):
                table.cell(i, j).text = cell

    doc.save(docx_path)
    print(f"Saved to {docx_path}")

convert_md_to_docx("test_results_grok.md", "test_results_grok.docx")
